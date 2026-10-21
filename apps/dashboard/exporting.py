"""Turn a test, writing task or speaking topic into a printable paper.

Every format is drawn from one list of elements, so the on-screen preview,
the Word file and the PDF cannot disagree about what a paper contains:

    ("title", text)                 ("heading", text)
    ("subtitle", text)              ("subheading", text, centered)
    ("para", text, style)           style: "", "label:A", "italic", "center"
    ("question", number, prompt, options)   options: [(label, text)]
    ("box", lines)                  bordered: an option pool, a cue card, a prompt
    ("image", path, url)            a local file, and its URL for the preview
    ("table", rows)                 the answer key
    ("pagebreak",)

Text may carry the app's **bold** markers; each renderer turns them into its
own bold. JSON is not a paper: it is the content itself, answers optional.

Layout follows the IELTS papers staff print from: A4, Times, bold "PART n"
and "Questions x-y" headings, blanks as "1 ______", options as "A) text".
"""

import io
import json
from dataclasses import dataclass
from xml.sax.saxutils import escape

from apps.common.text import split_bold
from apps.content.enums import BlockKind, QuestionType, Skill

FORMATS = ("pdf", "docx", "json")
CONTENT_TYPES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "json": "application/json",
}


@dataclass(frozen=True)
class Options:
    answers: bool = True
    explanations: bool = False
    transcript: bool = False
    model_answers: bool = False

    @classmethod
    def from_query(cls, query) -> "Options":
        # Checkboxes send nothing when unticked, so a submitted form is told
        # apart from a bare link by the presence of `format`.
        submitted = "format" in query

        def flag(name, default):
            return query.get(name) == "1" if submitted else default

        return cls(
            answers=flag("answers", True),
            explanations=flag("explanations", False),
            transcript=flag("transcript", False),
            model_answers=flag("model_answers", False),
        )


def _span(numbers) -> str:
    numbers = sorted(numbers)
    if not numbers:
        return ""
    if numbers[0] == numbers[-1]:
        return f"Question {numbers[0]}"
    return f"Questions {numbers[0]}\u2013{numbers[-1]}"


def _options(items) -> list[tuple[str, str]]:
    return [(o.label, o.text) for o in items]


def _answer_text(question) -> str:
    keys = list(question.answer_keys.all())
    shown = []
    for key in keys:
        label = key.option.label if key.option_id and key.option.label else ""
        shown.append(label or key.value)
    return " / ".join(dict.fromkeys(shown))


# --- builders -------------------------------------------------------------------


def test_queryset():
    from django.db.models import Prefetch

    from apps.content.models import Option, Question, Test

    return Test.objects.prefetch_related(
        "sections__blocks",
        "sections__groups",
        Prefetch(
            "sections__groups__questions",
            queryset=Question.objects.order_by("number").prefetch_related(
                Prefetch("options", queryset=Option.objects.order_by("order")),
                "answer_keys__option",
            ),
        ),
        Prefetch(
            "sections__groups__options",
            queryset=Option.objects.filter(question__isnull=True).order_by("order"),
        ),
    )


def test_paper(test, opts: Options) -> list[tuple]:
    reading = test.skill == Skill.READING
    out = [("title", f"IELTS {test.get_skill_display()}"), ("subtitle", test.title)]
    sections = sorted(test.sections.all(), key=lambda s: s.order)
    questions = []

    for index, section in enumerate(sections, start=1):
        groups = sorted(section.groups.all(), key=lambda g: g.order)
        numbers = [q.number for g in groups for q in g.questions.all()]
        blocks = sorted(section.blocks.all(), key=lambda b: b.order)

        if reading:
            if index > 1:
                out.append(("pagebreak",))
            out.append(("heading", f"READING PASSAGE {index}"))
            if numbers:
                out.append(
                    (
                        "para",
                        f"You should spend about 20 minutes on {_span(numbers)}, "
                        f"which are based on Reading Passage {index} below.",
                        "",
                    )
                )
            out.append(("subheading", section.title, True))
            out += _blocks(blocks)
        else:
            heading = f"PART {index}"
            if numbers:
                heading += f"    {_span(numbers)}"
            out.append(("heading", heading))

        for group in groups:
            group_questions = list(group.questions.all())
            questions += group_questions
            out.append(("subheading", _span([q.number for q in group_questions]), False))
            if group.instructions:
                out.append(("para", group.instructions, ""))
            if group.type == QuestionType.TFNG:
                out.append(("para", "TRUE / FALSE / NOT GIVEN", "italic"))
            if group.heading:
                out.append(("subheading", group.heading, True))
            pool = list(group.options.all())
            if pool:
                out.append(
                    ("box", [f"**{o.label}**  {o.text}" if o.label else o.text for o in pool])
                )
            for q in group_questions:
                out.append(("question", q.number, q.prompt, _options(q.options.all())))

    if opts.answers and questions:
        out += [("pagebreak",), ("heading", "ANSWER KEY")]
        out.append(("table", [(str(q.number), _answer_text(q)) for q in questions]))

    if opts.explanations and any(q.explanation for q in questions):
        out += [("pagebreak",), ("heading", "EXPLANATIONS")]
        for q in questions:
            if q.explanation:
                out.append(("para", f"**{q.number}  {_answer_text(q)}**", ""))
                out.append(("para", q.explanation, ""))

    if opts.transcript and not reading and any(s.blocks.all() for s in sections):
        out += [("pagebreak",), ("heading", "TRANSCRIPT")]
        for index, section in enumerate(sections, start=1):
            blocks = sorted(section.blocks.all(), key=lambda b: b.order)
            if blocks:
                out.append(("subheading", f"Part {index}", False))
                out += _blocks(blocks)
    return out


def _blocks(blocks) -> list[tuple]:
    out = []
    for block in blocks:
        if block.kind == BlockKind.RULE:
            continue
        if block.kind in (BlockKind.SECTION, BlockKind.HEADING):
            out.append(("subheading", block.text, block.kind == BlockKind.SECTION))
        elif block.text:
            out.append(("para", block.text, f"label:{block.label}" if block.label else ""))
    return out


def writing_paper(task, opts: Options) -> list[tuple]:
    one = task.task_number == 1
    out = [
        ("title", "IELTS Writing"),
        ("heading", f"WRITING TASK {task.task_number}"),
        ("para", f"You should spend about {task.suggested_time_minutes} minutes on this task.", ""),
    ]
    if not one:
        out.append(("para", "Write about the following topic:", ""))
    out.append(("box", [f"**{line}**" for line in task.prompt.split("\n\n") if line.strip()]))
    if task.image:
        try:
            out.append(("image", task.image.path, task.image.url))
        except NotImplementedError:  # storage without local paths
            pass
    if one:
        if task.type != "task1_letter":
            out.append(
                (
                    "para",
                    "Summarise the information by selecting and reporting the main features, "
                    "and make comparisons where relevant.",
                    "",
                )
            )
    else:
        out.append(
            (
                "para",
                "Give reasons for your answer and include any relevant examples from your "
                "own knowledge or experience.",
                "",
            )
        )
    out.append(("para", f"Write at least {task.target_words} words.", ""))

    answers = list(task.model_answers.order_by("order")) if opts.model_answers else []
    if answers:
        out += [("pagebreak",), ("heading", "MODEL ANSWER" + ("S" if len(answers) > 1 else ""))]
        for answer in answers:
            if answer.band is not None:
                out.append(("subheading", f"Band {answer.band}", False))
            out += [("para", p, "") for p in answer.paragraphs if p.strip()]
    return out


SPEAKING_PARTS = {
    1: "PART 1 — INTRODUCTION AND INTERVIEW",
    2: "PART 2 — LONG TURN",
    3: "PART 3 — DISCUSSION",
}


def speaking_paper(topic, opts: Options) -> list[tuple]:
    from apps.speaking.models import SpeakingItemKind

    items = sorted(topic.items.all(), key=lambda i: (i.part, i.kind, i.order))

    def texts(part, kind):
        return [i.text for i in items if i.part == part and i.kind == kind]

    out = [("title", "IELTS Speaking"), ("subtitle", topic.title)]
    if texts(1, SpeakingItemKind.QUESTION):
        out.append(("heading", SPEAKING_PARTS[1]))
        out += [
            ("question", n, t, []) for n, t in enumerate(texts(1, SpeakingItemKind.QUESTION), 1)
        ]
    cue = getattr(topic, "cue_card", None)
    if cue:
        out.append(("heading", SPEAKING_PARTS[2]))
        out.append(
            (
                "para",
                "You will have to talk about the topic for one to two minutes. "
                "You have one minute to think about what you are going to say.",
                "italic",
            )
        )
        bullets = texts(2, SpeakingItemKind.BULLET)
        out.append(("box", [f"**{cue.title}**", "You should say:", *[f"•  {b}" for b in bullets]]))
    if texts(3, SpeakingItemKind.QUESTION):
        out.append(("heading", SPEAKING_PARTS[3]))
        out += [
            ("question", n, t, []) for n, t in enumerate(texts(3, SpeakingItemKind.QUESTION), 1)
        ]
    return out


# --- JSON --------------------------------------------------------------------


def test_json(test, opts: Options) -> dict:
    sections = []
    for section in sorted(test.sections.all(), key=lambda s: s.order):
        groups = []
        for group in sorted(section.groups.all(), key=lambda g: g.order):
            questions = []
            for q in group.questions.all():
                item = {
                    "number": q.number,
                    "prompt": q.prompt,
                    "options": _options(q.options.all()),
                }
                if opts.answers:
                    item["answers"] = [k.value for k in q.answer_keys.all()]
                if opts.explanations and q.explanation:
                    item["explanation"] = q.explanation
                questions.append(item)
            groups.append(
                {
                    "type": group.type,
                    "heading": group.heading,
                    "instructions": group.instructions,
                    "options": _options(group.options.all()),
                    "questions": questions,
                }
            )
        data = {"title": section.title, "instructions": section.instructions, "groups": groups}
        if test.skill == Skill.READING or opts.transcript:
            key = "label" if test.skill == Skill.READING else "speaker"
            data["blocks"] = [
                {"kind": b.kind, key: b.label, "text": b.text}
                for b in sorted(section.blocks.all(), key=lambda b: b.order)
            ]
        if section.audio_id:
            data["audio"] = {
                "url": section.audio.file.url,
                "startMs": section.audio_start_ms,
                "endMs": section.audio_end_ms,
            }
        sections.append(data)
    return {
        "id": test.slug,
        "skill": test.skill,
        "title": test.title,
        "description": test.description,
        "timeLimitMinutes": test.time_limit_minutes,
        "pack": test.pack.title if test.pack_id else None,
        "sections": sections,
    }


def writing_json(task, opts: Options) -> dict:
    data = {
        "id": task.slug,
        "taskNumber": task.task_number,
        "type": task.type,
        "prompt": task.prompt,
        "image": task.image.url if task.image else None,
        "suggestedTimeMinutes": task.suggested_time_minutes,
        "targetWords": task.target_words,
        "pack": task.pack.title if task.pack_id else None,
    }
    if opts.model_answers:
        data["modelAnswers"] = [
            {"band": str(a.band) if a.band is not None else None, "body": a.body}
            for a in task.model_answers.order_by("order")
        ]
    return data


def speaking_json(topic, opts: Options) -> dict:
    from apps.speaking.publishing import build_topic_payload

    return build_topic_payload(topic)


# --- renderers ----------------------------------------------------------------


def to_json(data: dict) -> bytes:
    return json.dumps(data, ensure_ascii=False, indent=2).encode()


def html_elements(elements: list[tuple]) -> list[dict]:
    """Elements as dicts with **bold** split, for the preview template."""
    out = []
    for el in elements:
        kind = el[0]
        if kind in ("title", "subtitle", "heading"):
            out.append({"kind": kind, "text": el[1]})
        elif kind == "subheading":
            out.append({"kind": kind, "text": el[1], "center": el[2]})
        elif kind == "para":
            style = el[2]
            label = style[6:] if style.startswith("label:") else ""
            out.append({"kind": kind, "parts": split_bold(el[1]), "label": label, "style": style})
        elif kind == "question":
            out.append(
                {"kind": kind, "number": el[1], "parts": split_bold(el[2]), "options": el[3]}
            )
        elif kind == "box":
            out.append({"kind": kind, "lines": [split_bold(line) for line in el[1]]})
        elif kind == "image":
            out.append({"kind": kind, "url": el[2]})
        elif kind == "table":
            rows = el[1]
            half = (len(rows) + 1) // 2
            pairs = [
                (rows[i], rows[i + half] if i + half < len(rows) else ("", "")) for i in range(half)
            ]
            out.append({"kind": kind, "rows": pairs})
        else:
            out.append({"kind": kind})
    return out


def _png_bytes(path: str) -> tuple[bytes, int, int] | None:
    """Any image Pillow reads (avif included), as PNG for Word and PDF."""
    from PIL import Image

    try:
        with Image.open(path) as image:
            image = image.convert("RGB")
            buffer = io.BytesIO()
            image.save(buffer, format="PNG")
            return buffer.getvalue(), image.width, image.height
    except (OSError, ValueError):
        return None


def _option_line(options) -> str:
    labels = [label for label, _ in options]
    if not any(labels):
        return " / ".join(text for _, text in options)
    return ""


def to_docx(elements: list[tuple]) -> bytes:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    doc = Document()
    section = doc.sections[0]
    section.page_height, section.page_width = Cm(29.7), Cm(21)
    section.top_margin = section.bottom_margin = Cm(1.27)
    section.left_margin, section.right_margin = Cm(2), Cm(1.27)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.paragraph_format.space_after = Pt(4)

    def runs(paragraph, text, bold=False, italic=False, size=None):
        for part in split_bold(text or ""):
            run = paragraph.add_run(part["text"])
            run.bold = bold or part["bold"]
            run.italic = italic
            if size:
                run.font.size = Pt(size)

    def bottom_rule(paragraph):
        p_pr = paragraph._p.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for attr, value in (
            ("w:val", "single"),
            ("w:sz", "6"),
            ("w:space", "4"),
            ("w:color", "000000"),
        ):
            bottom.set(qn(attr), value)
        borders.append(bottom)
        p_pr.append(borders)

    for el in elements:
        kind = el[0]
        if kind == "title":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            runs(p, el[1], bold=True, size=16)
        elif kind == "subtitle":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(12)
            runs(p, el[1], size=13)
        elif kind == "heading":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.keep_with_next = True
            runs(p, el[1], bold=True, size=15)
            bottom_rule(p)
        elif kind == "subheading":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.keep_with_next = True
            if el[2]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            runs(p, el[1], bold=True)
        elif kind == "para":
            style = el[2]
            p = doc.add_paragraph()
            if style.startswith("label:"):
                runs(p, style[6:] + "  ", bold=True)
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if style == "center":
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            runs(p, el[1], italic=style == "italic")
        elif kind == "question":
            _, number, prompt, options = el
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            runs(p, f"{number}  ", bold=True)
            runs(p, prompt)
            inline = _option_line(options)
            if inline:
                runs(doc.add_paragraph(), inline, italic=True)
            elif options:
                for label, text in options:
                    o = doc.add_paragraph()
                    o.paragraph_format.left_indent = Cm(0.8)
                    o.paragraph_format.space_after = Pt(1)
                    runs(o, f"{label}) ", bold=True)
                    runs(o, text)
        elif kind == "box":
            table = doc.add_table(rows=1, cols=1)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = table.rows[0].cells[0]
            cell.paragraphs[0].text = ""
            for i, line in enumerate(el[1]):
                p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
                runs(p, line)
            doc.add_paragraph()
        elif kind == "image":
            png = _png_bytes(el[1])
            if png:
                data, width, _height = png
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(io.BytesIO(data), width=Cm(min(16, width / 37.8)))
        elif kind == "table":
            rows = el[1]
            half = (len(rows) + 1) // 2
            table = doc.add_table(rows=half, cols=4)
            table.style = "Table Grid"
            for i in range(half):
                for col, pair in enumerate(
                    (rows[i], rows[i + half] if i + half < len(rows) else ("", ""))
                ):
                    number_cell, answer_cell = table.rows[i].cells[col * 2 : col * 2 + 2]
                    number_cell.paragraphs[0].text = ""
                    runs(number_cell.paragraphs[0], pair[0], bold=True)
                    answer_cell.paragraphs[0].text = pair[1]
        elif kind == "pagebreak":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def to_pdf(elements: list[tuple], title: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        Image,
        KeepTogether,
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    base = ParagraphStyle("base", fontName="Times-Roman", fontSize=11.5, leading=15, spaceAfter=4)
    styles = {
        "title": ParagraphStyle(
            "title", base, fontName="Times-Bold", fontSize=17, alignment=TA_CENTER, leading=21
        ),
        "subtitle": ParagraphStyle(
            "subtitle", base, fontSize=13, alignment=TA_CENTER, spaceAfter=12
        ),
        "heading": ParagraphStyle(
            "heading", base, fontName="Times-Bold", fontSize=15, spaceBefore=12, leading=19
        ),
        "sub": ParagraphStyle("sub", base, fontName="Times-Bold", spaceBefore=8),
        "subc": ParagraphStyle(
            "subc", base, fontName="Times-Bold", spaceBefore=8, alignment=TA_CENTER
        ),
        "body": base,
        "just": ParagraphStyle("just", base, alignment=TA_JUSTIFY),
        "italic": ParagraphStyle("italic", base, fontName="Times-Italic"),
        "center": ParagraphStyle("center", base, alignment=TA_CENTER),
        "option": ParagraphStyle("option", base, leftIndent=0.8 * cm, spaceAfter=1),
    }
    # The frame keeps 6pt of padding on each side; a table as wide as the
    # margins would poke out past the text by that much.
    width = A4[0] - 3.27 * cm - 12

    def markup(text, bold=False):
        parts = []
        for part in split_bold(text or ""):
            chunk = escape(part["text"])
            parts.append(f"<b>{chunk}</b>" if part["bold"] and not bold else chunk)
        # Runs of spaces are deliberate ("PART 1    Questions 1-10"); PDF text collapses them.
        out = "".join(parts).replace("  ", "&nbsp; ")
        return f"<b>{out}</b>" if bold else out

    def rule(flowable):
        table = Table([[flowable]], colWidths=[width])
        table.setStyle(
            TableStyle(
                [
                    ("LINEBELOW", (0, 0), (-1, -1), 0.8, colors.black),
                    ("LEFTPADDING", (0, 0), (-1, -1), 0),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ]
            )
        )
        return table

    story = []
    for el in elements:
        kind = el[0]
        if kind == "title":
            story.append(Paragraph(markup(el[1]), styles["title"]))
        elif kind == "subtitle":
            story.append(Paragraph(markup(el[1]), styles["subtitle"]))
        elif kind == "heading":
            story += [Spacer(1, 8), rule(Paragraph(markup(el[1]), styles["heading"])), Spacer(1, 4)]
        elif kind == "subheading":
            story.append(Paragraph(markup(el[1]), styles["subc" if el[2] else "sub"]))
        elif kind == "para":
            style = el[2]
            if style.startswith("label:"):
                story.append(
                    Paragraph(
                        f"<b>{escape(style[6:])}</b>&nbsp;&nbsp;{markup(el[1])}", styles["just"]
                    )
                )
            else:
                story.append(Paragraph(markup(el[1]), styles.get(style or "body", styles["body"])))
        elif kind == "question":
            _, number, prompt, options = el
            block = [Paragraph(f"<b>{number}</b>&nbsp;&nbsp;{markup(prompt)}", styles["body"])]
            inline = _option_line(options)
            if inline:
                block.append(Paragraph(markup(inline), styles["italic"]))
            else:
                block += [
                    Paragraph(f"<b>{escape(label)})</b> {markup(text)}", styles["option"])
                    for label, text in options
                ]
            story.append(KeepTogether(block))
        elif kind == "box":
            cell = [Paragraph(markup(line), styles["body"]) for line in el[1]]
            table = Table([[cell]], colWidths=[width])
            table.setStyle(
                TableStyle(
                    [
                        ("BOX", (0, 0), (-1, -1), 0.8, colors.black),
                        ("LEFTPADDING", (0, 0), (-1, -1), 8),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                        ("TOPPADDING", (0, 0), (-1, -1), 6),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                    ]
                )
            )
            story += [table, Spacer(1, 8)]
        elif kind == "image":
            png = _png_bytes(el[1])
            if png:
                data, w, h = png
                shown = min(width, 16 * cm)
                story += [Image(io.BytesIO(data), width=shown, height=shown * h / w), Spacer(1, 6)]
        elif kind == "table":
            rows = el[1]
            half = (len(rows) + 1) // 2
            data = []
            for i in range(half):
                right = rows[i + half] if i + half < len(rows) else ("", "")
                data.append(
                    [
                        Paragraph(f"<b>{escape(rows[i][0])}</b>", styles["body"]),
                        Paragraph(escape(rows[i][1]), styles["body"]),
                        Paragraph(f"<b>{escape(right[0])}</b>", styles["body"]),
                        Paragraph(escape(right[1]), styles["body"]),
                    ]
                )
            if data:
                table = Table(data, colWidths=[width * 0.1, width * 0.4, width * 0.1, width * 0.4])
                table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
                story.append(table)
        elif kind == "pagebreak":
            story.append(PageBreak())

    buffer = io.BytesIO()
    SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=2 * cm,
        rightMargin=1.27 * cm,
        topMargin=1.27 * cm,
        bottomMargin=1.27 * cm,
        title=title,
    ).build(story)
    return buffer.getvalue()
